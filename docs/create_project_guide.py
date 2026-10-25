from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from docx.enum.style import WD_STYLE_TYPE
from pathlib import Path

OUT = Path(__file__).parent / 'Guia_del_proyecto_BuscadorDocs.docx'
GREEN, PALE, GRAY, TEXT = '1F5C4A', 'EAF2EC', 'D9DFD5', '1E302B'
doc = Document()
sec = doc.sections[0]
sec.top_margin = Cm(2.0); sec.bottom_margin = Cm(1.8)
sec.left_margin = Cm(2.2); sec.right_margin = Cm(2.2)
styles = doc.styles
styles['Normal'].font.name = 'Aptos'
styles['Normal']._element.rPr.rFonts.set(qn('w:eastAsia'), 'Aptos')
styles['Normal'].font.size = Pt(10.5); styles['Normal'].font.color.rgb = RGBColor.from_string(TEXT)
styles['Normal'].paragraph_format.space_after = Pt(7); styles['Normal'].paragraph_format.line_spacing = 1.12
for name, size in [('Title', 25), ('Heading 1', 17), ('Heading 2', 13), ('Heading 3', 11)]:
    styles[name].font.name = 'Aptos Display' if name != 'Heading 3' else 'Aptos'
    styles[name]._element.rPr.rFonts.set(qn('w:eastAsia'), styles[name].font.name)
    styles[name].font.size = Pt(size); styles[name].font.bold = True; styles[name].font.color.rgb = RGBColor(0, 0, 0)
    styles[name].paragraph_format.space_before = Pt(13 if name == 'Heading 1' else 9)
    styles[name].paragraph_format.space_after = Pt(5)
try: styles.add_style('Small Note', WD_STYLE_TYPE.PARAGRAPH)
except ValueError: pass
styles['Small Note'].font.name = 'Aptos'; styles['Small Note'].font.size = Pt(9); styles['Small Note'].font.color.rgb = RGBColor.from_string('63726B')

def shade(cell, color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn('w:shd')) or OxmlElement('w:shd')
    shd.set(qn('w:fill'), color)
    if shd.getparent() is None: tcPr.append(shd)

def borders(cell, color=GRAY):
    tcPr = cell._tc.get_or_add_tcPr()
    b = tcPr.first_child_found_in('w:tcBorders')
    if b is None: b = OxmlElement('w:tcBorders'); tcPr.append(b)
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        node = b.find(qn('w:' + edge)) or OxmlElement('w:' + edge)
        node.set(qn('w:val'), 'single'); node.set(qn('w:sz'), '4'); node.set(qn('w:color'), color)
        if node.getparent() is None: b.append(node)

def set_cell_text(cell, text, bold=False, color=TEXT, size=9.5):
    cell.text = ''
    p = cell.paragraphs[0]; p.paragraph_format.space_after = Pt(3)
    run = p.add_run(str(text)); run.bold = bold; run.font.name = 'Aptos'; run.font.size = Pt(size); run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

def table(headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers)); t.alignment = WD_TABLE_ALIGNMENT.CENTER; t.autofit = True
    trPr = t.rows[0]._tr.get_or_add_trPr()
    tbl_header = OxmlElement('w:tblHeader'); tbl_header.set(qn('w:val'), 'true'); trPr.append(tbl_header)
    for i, h in enumerate(headers):
        set_cell_text(t.rows[0].cells[i], h, True, 'FFFFFF', 9); shade(t.rows[0].cells[i], GREEN); borders(t.rows[0].cells[i])
    for r, row in enumerate(rows):
        cells = t.add_row().cells
        for i, value in enumerate(row):
            set_cell_text(cells[i], value, False, TEXT, 9)
            if r % 2 == 1: shade(cells[i], 'F5F7F3')
            borders(cells[i])
    if widths:
        for row in t.rows:
            for i, width in enumerate(widths): row.cells[i].width = Cm(width)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return t

def bullet(text):
    p = doc.add_paragraph(style='List Bullet'); p.add_run(text); return p

def code(text):
    p = doc.add_paragraph(); p.paragraph_format.left_indent = Cm(.4); p.paragraph_format.right_indent = Cm(.4)
    p.paragraph_format.space_before = Pt(3); p.paragraph_format.space_after = Pt(7)
    run = p.add_run(text); run.font.name = 'Consolas'; run._element.rPr.rFonts.set(qn('w:eastAsia'), 'Consolas')
    run.font.size = Pt(8.5); run.font.color.rgb = RGBColor.from_string('2E5046'); return p

# Portada
p = doc.add_paragraph(style='Title'); p.add_run('Guía del proyecto BuscadorDocs')
p = doc.add_paragraph(); r = p.add_run('Sistema de búsqueda, minería bibliográfica y análisis reproducible')
r.bold = True; r.font.size = Pt(14); r.font.color.rgb = RGBColor.from_string(GREEN)
doc.add_paragraph('Documento de explicación funcional y técnica para comprender el propósito del proyecto, el uso de cada sección y el recorrido de los datos.', style='Subtitle')
doc.add_paragraph('Versión del documento: 11 de octubre de 2026', style='Small Note')
doc.add_heading('Qué debes entender al terminar', level=1)
doc.add_paragraph('BuscadorDocs convierte una pregunta temática en un recorrido de investigación: recupera documentos, los organiza, extrae señales del corpus, genera un informe en español con referencias y código R, y permite guardar el resultado en una biblioteca reutilizable. La aplicación está pensada para explorar cualquier tema con trazabilidad suficiente para revisar de dónde sale cada afirmación.')
table(['Etapa', 'Resultado'], [
    ('1. Buscar', 'El sistema interpreta un tema y consulta la evidencia bibliográfica disponible.'),
    ('2. Explorar', 'Muestra palabras frecuentes, años, fuentes y documentos usados como datos.'),
    ('3. Investigar', 'Construye un informe estructurado con definición, hallazgos, aplicaciones y R.'),
    ('4. Experimentar', 'Permite trabajar con CSV sintético o con datos extraídos de documentos recuperados.'),
    ('5. Guardar', 'Almacena el LaTeX en Supabase para consultarlo como tema y previsualizarlo.'),
], [3.3, 12.2])

doc.add_page_break()
doc.add_heading('1 Propósito del proyecto', level=1)
doc.add_paragraph('El proyecto funciona como un centro de investigación asistida para minería de datos y vigilancia bibliográfica. Su objetivo no es solamente listar resultados de Google o OpenAlex: busca ayudar a pasar de una palabra o pregunta inicial a un corpus manejable, una lectura rápida del tema, un análisis de datos y un informe que pueda revisarse y reutilizarse.')
doc.add_paragraph('El usuario escribe un tema como programación, datos, inteligencia artificial, Scrum o predicción de ventas. A partir de esa entrada, la aplicación conecta fuentes bibliográficas, reglas de selección, extracción de términos, ejemplos en R y almacenamiento persistente. La interfaz separa la exploración rápida, la investigación, el laboratorio y la biblioteca para que el contenido no aparezca como una única cascada.')
doc.add_heading('Valor para una evaluación o licitación', level=2)
for text in ['Convierte consultas abiertas en un flujo demostrable y repetible.', 'Hace visible la trazabilidad: documentos, puntajes, decisiones, años, fuentes y datos faltantes.', 'Relaciona vigilancia bibliográfica con minería de texto y experimentación en R.', 'Permite ampliar el sistema a nuevos temas sin diseñar una pantalla por tema.', 'Conserva informes en Supabase y los vuelve a integrar en la biblioteca como temas navegables.']: bullet(text)

doc.add_heading('2 Arquitectura general', level=1)
doc.add_paragraph('La aplicación está dividida en un frontend React, un backend Node.js y servicios externos. El frontend gestiona navegación, búsqueda, componentes visuales y acciones del usuario. El backend concentra la lógica de recuperación, normalización, análisis, generación de informes y conexión con Supabase. OpenAlex entrega metadatos bibliográficos y OpenRouter genera el texto estructurado del informe cuando se solicita el modo de IA.')
table(['Capa', 'Responsabilidad', 'Elementos principales'], [
    ('Interfaz', 'Presentar y organizar el trabajo de investigación.', 'React, Vite, componentes de búsqueda, biblioteca, laboratorio, dashboard y chat.'),
    ('API local', 'Orquestar las operaciones y proteger las claves.', 'Node.js, rutas /api, validación, normalización y manejo de errores.'),
    ('Evidencia', 'Recuperar y puntuar documentos.', 'OpenAlex, deduplicación por DOI o identificador, años, citas y resúmenes.'),
    ('IA', 'Redactar el informe completo en español.', 'OpenRouter con modelos gratuitos, salida JSON estructurada y reintentos.'),
    ('Persistencia', 'Guardar informes y volverlos consultables.', 'Supabase Storage, archivos .tex y carga de temas guardados.'),
    ('Análisis', 'Explorar datos tabulares y texto documental.', 'CSV, perfil de datos, frecuencia, TF-IDF, evolución temporal y relaciones.'),
], [3.0, 5.3, 7.2])
doc.add_heading('Recorrido resumido', level=2)
code('Tema → OpenAlex → deduplicación y puntaje → corpus seleccionado → minería del corpus → informe JSON → revisión → LaTeX → Supabase → biblioteca')

doc.add_heading('3 Inicio y navegación', level=1)
doc.add_paragraph('La pantalla inicial funciona como punto de entrada. Resume cuántos temas, áreas e informes guardados existen y ofrece accesos al laboratorio y a la biblioteca. El menú lateral agrupa temas por áreas como Datos, Desarrollo, Bases de datos, Estadística, IA, Machine learning y Minería de datos.')
doc.add_paragraph('Las áreas son menús desplegables. Un tema puede pertenecer a varias áreas, pero la interfaz evita mostrar una lista plana interminable y cada tema se mantiene como una única entidad de navegación. Al seleccionar un tema se abre su ficha con definición, utilidad, ejemplo en R, referencias y acciones de investigación.')
table(['Elemento', 'Para qué sirve'], [
    ('Inicio', 'Explica el propósito y dirige al flujo de investigación.'),
    ('Laboratorio', 'Carga o genera CSV y ejecuta análisis.'),
    ('Chat rápido', 'Responde preguntas concisas sobre minería de datos y R.'),
    ('Biblioteca', 'Organiza temas locales y documentos guardados.'),
    ('Dashboard', 'Resume señales del corpus de un tema.'),
], [3.4, 12.1])

doc.add_page_break()
doc.add_heading('4 Búsqueda y exploración de evidencia', level=1)
doc.add_paragraph('La búsqueda rápida se activa al escribir un tema. La aplicación consulta OpenAlex y recupera metadatos como título, autores, año, fuente, DOI, citas y resumen. Después elimina duplicados y calcula un puntaje compuesto para priorizar los documentos que mejor coinciden con el tema.')
doc.add_heading('Qué aparece en la sección de coincidencias', level=2)
for text in ['Un encabezado que indica que se está explorando la evidencia del tema.', 'Un panorama con los términos más frecuentes y su número de apariciones.', 'Una línea temporal que permite ver los años con documentos y los años sin cobertura.', 'Una vista de documentos con título, autores, año, fuente, citas y enlace.', 'Un panel de código R que muestra cómo explorar el corpus documental.']: bullet(text)
doc.add_heading('Cómo leer las palabras clave', level=2)
doc.add_paragraph('Las palabras clave se extraen de títulos y resúmenes después de filtrar palabras vacías y el propio tema. El número indica frecuencia observada en el corpus, no importancia científica ni causalidad. Por eso la interfaz incluye la aclaración de que una frecuencia alta debe interpretarse junto con los documentos y sus resúmenes.')
doc.add_heading('Puntaje y decisiones', level=2)
doc.add_paragraph('El puntaje prioriza coincidencia temática, unidad del tema, calidad de la fuente, presencia de señales metodológicas, actualidad y aplicabilidad. La decisión “incluir” identifica evidencia fuerte para el corpus seleccionado; “respaldo” conserva documentos útiles como contexto; “excluir” indica que el documento no entra en la selección principal.')

doc.add_heading('5 Agente investigador', level=1)
doc.add_paragraph('El agente investigador recibe el tema y el corpus recuperado. Su trabajo es redactar un informe completo en español a partir de esa evidencia. El modelo no debe buscar por su cuenta ni inventar referencias: las referencias finales se mantienen vinculadas al corpus que preparó el backend.')
doc.add_heading('Contenido del informe', level=2)
table(['Sección', 'Contenido'], [
    ('Alcance y desambiguación', 'Aclara qué sentido del tema se estudia y qué queda fuera.'),
    ('Definición', 'Explica qué es el tema y lo relaciona con la evidencia.'),
    ('Minería del corpus', 'Resume documentos, puntajes, años, fuentes y términos.'),
    ('Hallazgos', 'Presenta observaciones sustantivas con citas del corpus.'),
    ('Oportunidades', 'Describe aplicaciones concretas respaldadas por documentos.'),
    ('Ejemplo en R', 'Incluye código ejecutable y una explicación de su propósito.'),
    ('Referencias y trazabilidad', 'Registra tipo, fuente, año, decisión, puntaje y motivo.'),
    ('Limitaciones', 'Declara vacíos, cobertura temporal y verificaciones pendientes.'),
], [4.1, 11.4])
doc.add_paragraph('El sistema valida que existan campos esenciales, rechaza respuestas incompletas y pide JSON estructurado. También puede reintentar con otro modelo gratuito cuando un proveedor está saturado o devuelve una respuesta vacía. La interfaz muestra un error comprensible en lugar de presentar un informe parcial como si fuera válido.')

doc.add_heading('6 Biblioteca e informes guardados', level=1)
doc.add_paragraph('Cuando el usuario guarda una investigación, el backend genera un archivo LaTeX ampliado y lo sube a Supabase Storage. El archivo contiene más información que la ficha resumida: trazabilidad, pipeline, definición, alcance, aplicaciones, hallazgos, conceptos clave, código R, referencias, auditoría editorial, registro de cambios, datos faltantes y nota metodológica.')
doc.add_paragraph('Al iniciar la aplicación, los archivos guardados se cargan desde Supabase y se convierten en temas navegables. Esto permite buscar nuevamente un informe, abrir su ficha y previsualizarlo. El frontend limpia entidades HTML, elimina encabezados duplicados y evita mostrar varias veces un marcador de falta de evidencia.')
table(['Acción', 'Resultado'], [
    ('Descargar Markdown', 'Obtiene una versión portable para lectura o entrega.'),
    ('Previsualizar', 'Abre el informe guardado en la vista de documento o PDF del proyecto.'),
    ('Guardar informe', 'Genera el .tex y lo almacena en Supabase.'),
    ('Consultar fuente', 'Abre el DOI o identificador de OpenAlex del documento.'),
], [4.0, 11.5])

doc.add_page_break()
doc.add_heading('7 Laboratorio de minería de datos', level=1)
doc.add_paragraph('El laboratorio es el espacio experimental. Permite elegir entre dos fuentes de datos para que el análisis sea útil tanto para una demostración como para un escenario de investigación real.')
table(['Fuente elegida', 'Qué hace', 'Cuándo usarla'], [
    ('Datos sintéticos', 'Genera un CSV de prueba relacionado con el tema elegido.', 'Para probar rápidamente métodos, gráficos y métricas.'),
    ('Documentos recuperados', 'Construye un CSV con títulos, autores, años, fuentes, citas, puntajes, decisiones, DOI y resúmenes.', 'Para realizar minería textual sobre evidencia bibliográfica real.'),
], [3.3, 7.0, 5.2])
doc.add_heading('Análisis disponibles', level=2)
for text in ['Perfil del dataset: columnas, tipos, faltantes, valores únicos y calidad básica.', 'Estadística descriptiva: medidas resumen para variables numéricas.', 'Frecuencias y distribución: identifica valores frecuentes y concentraciones.', 'Clasificación y regresión: ofrece modelos de ejemplo cuando el dataset lo permite.', 'Minería textual: calcula términos frecuentes, frecuencia documental, TF-IDF, años, fuentes y coocurrencias.']: bullet(text)
doc.add_paragraph('La opción de documentos recuperados es la más alineada con el objetivo del proyecto: transforma la bibliografía en un dataset analizable y permite estudiar no solo cuántos documentos hay, sino también qué conceptos aparecen, cómo se distribuyen en el tiempo y qué fuentes concentran la evidencia.')

doc.add_heading('8 Dashboard', level=1)
doc.add_paragraph('El dashboard ofrece una lectura resumida del corpus de un tema guardado. Presenta métricas de volumen, términos frecuentes, evolución documental, fuentes principales y señales de cobertura. Su función es apoyar una decisión rápida: saber si un tema tiene suficiente evidencia, si está concentrado en pocos años o si requiere ampliar la búsqueda.')
table(['Indicador', 'Interpretación'], [
    ('Documentos', 'Volumen de evidencia disponible en el corpus.'),
    ('Términos', 'Señales léxicas observadas en títulos y resúmenes.'),
    ('Años', 'Cobertura temporal y posibles vacíos.'),
    ('Fuentes', 'Repositorios, revistas o editoriales con más presencia.'),
    ('Calidad y faltantes', 'Primera revisión de completitud de los metadatos.'),
], [3.3, 12.2])

doc.add_heading('9 Chat rápido', level=1)
doc.add_paragraph('El chatbot responde preguntas breves sobre minería de datos, análisis de datos, aprendizaje automático aplicado, estadística, bases de datos analíticas, visualización y ejemplos en R. Responde en español y limita su alcance para no convertirse en un asistente general sin relación con el proyecto.')
doc.add_paragraph('Cuando el usuario pide R, el chatbot debe entregar código ejecutable, comentarios en español y una explicación corta. Si la respuesta del proveedor incluye razonamiento interno o una estructura no válida, el backend la limpia o entrega una respuesta de respaldo.')

doc.add_heading('10 Flujo de uso recomendado', level=1)
for text in ['Abre Inicio y revisa las áreas y los informes guardados.', 'Escribe un tema concreto en el buscador.', 'Lee el panorama: documentos, términos y años.', 'Abre Documentos para revisar los títulos que sustentan la minería.', 'Pulsa Generar informe y selecciona OpenRouter gratis para producir el informe en español.', 'Revisa definición, hallazgos, aplicaciones, ejemplo R y referencias.', 'Usa el Laboratorio con documentos recuperados para profundizar en el corpus.', 'Guarda el informe cuando hayas revisado las fuentes.', 'Consulta el informe desde Biblioteca y usa Dashboard para comparar temas.']:
    p = doc.add_paragraph(style='List Number'); p.add_run(text)

doc.add_heading('11 Rutas y archivos principales', level=1)
table(['Archivo o ruta', 'Función'], [
    ('frontend/src/main.jsx', 'Componentes principales, navegación, búsqueda, biblioteca y laboratorio.'),
    ('frontend/src/ResearchPanels.jsx', 'Paneles visuales de evidencia e informe.'),
    ('frontend/src/research.css', 'Estilos de la exploración bibliográfica y lectura del informe.'),
    ('frontend/src/styles.css', 'Sistema visual general de la aplicación.'),
    ('backend/server.mjs', 'Servidor HTTP, rutas API, carga de configuración y orquestación.'),
    ('backend/agent.mjs', 'Recuperación OpenAlex, puntuación, minería y generación LaTeX.'),
    ('backend/openrouter-report.mjs', 'Solicitud a OpenRouter, JSON estructurado, validación y reintentos.'),
    ('backend/analysis.mjs', 'Perfil y análisis de CSV, incluido text mining.'),
    ('backend/supabase.mjs', 'Listado y carga de archivos LaTeX en Supabase Storage.'),
    ('shared/unidad-1/latex', 'Base documental local que alimenta temas iniciales.'),
], [6.2, 9.3])

doc.add_heading('12 Consideraciones y mejoras futuras', level=1)
for text in ['Verificar manualmente textos completos, revisión por pares y DOI antes de usar un informe como entrega académica final.', 'Agregar filtros por tipo de documento, idioma, año y área de aplicación.', 'Guardar versiones del informe para comparar cómo cambia un tema con nuevas búsquedas.', 'Incluir exportación a PDF final compilado desde LaTeX cuando el entorno de compilación esté disponible.', 'Agregar comparación entre dos temas en el dashboard mediante términos compartidos y diferencias de cobertura.', 'Incorporar una vista de calidad del corpus que destaque metadatos incompletos y fuentes poco confiables.']: bullet(text)
doc.add_paragraph('Estas mejoras mantienen la idea central del proyecto: una pregunta temática debe convertirse en evidencia organizada, datos analizables y un resultado trazable.', style='Small Note')

for section in doc.sections:
    footer = section.footer.paragraphs[0]; footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = footer.add_run('BuscadorDocs · Guía del proyecto'); run.font.size = Pt(8); run.font.color.rgb = RGBColor.from_string('63726B')
props = doc.core_properties
props.title = 'Guía del proyecto BuscadorDocs'; props.subject = 'Explicación funcional y técnica del sistema'
props.author = 'BuscadorDocs'; props.keywords = 'minería de datos, vigilancia bibliográfica, OpenAlex, OpenRouter, Supabase, R'
doc.save(OUT)
print(OUT)

